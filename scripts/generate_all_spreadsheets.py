import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FACILITIES = [
    {
        "code": "hml",
        "abbr": "HML",
        "name": "Hillside Medical Lodge",
        "city": "Beeville, TX",
        "shootDate": "Thursday, November 5, 2026",
        "shortDate": "11/5",
        "dayOfWeek": "Thursday",
        "loungeName": "Hillside Fireside Solarium",
    },
    {
        "code": "wnr",
        "abbr": "WNR",
        "name": "Whitney Nursing & Rehabilitation",
        "city": "Whitney, TX",
        "shootDate": "Friday, November 6, 2026",
        "shortDate": "11/6",
        "dayOfWeek": "Friday",
        "loungeName": "Whitney Heritage Community Room",
    },
    {
        "code": "cml",
        "abbr": "CML",
        "name": "Cheyenne Medical Lodge",
        "city": "Colorado City, TX",
        "shootDate": "Monday, November 9 & Tuesday, November 10, 2026",
        "shortDate": "11/9 & 11/10",
        "dayOfWeek": "Monday & Tuesday (2-Day Session)",
        "dates": [
            {"label": "Day 1 (Nov 9)", "dateStr": "Monday, November 9, 2026", "shortDate": "11/9"},
            {"label": "Day 2 (Nov 10)", "dateStr": "Tuesday, November 10, 2026", "shortDate": "11/10"},
        ],
        "loungeName": "Cheyenne Grand Prairie Staging Room",
    },
    {
        "code": "pml",
        "abbr": "PML",
        "name": "Princeton Medical Lodge",
        "city": "Princeton, TX",
        "shootDate": "Thursday, November 12, 2026",
        "shortDate": "11/12",
        "dayOfWeek": "Thursday",
        "loungeName": "Princeton Courtyard Pavilion",
    },
    {
        "code": "fhr",
        "abbr": "FHR",
        "name": "Farmersville Health & Rehabilitation",
        "city": "Farmersville, TX",
        "shootDate": "Friday, November 13, 2026",
        "shortDate": "11/13",
        "dayOfWeek": "Friday",
        "loungeName": "Farmersville Evergreen Great Room",
    },
    {
        "code": "lml",
        "abbr": "LML",
        "name": "Lexington Medical Lodge",
        "city": "Farmersville, TX",
        "shootDate": "Tuesday, November 17, 2026",
        "shortDate": "11/17",
        "dayOfWeek": "Tuesday",
        "loungeName": "Lexington Magnolia Activity Atrium",
    },
    {
        "code": "tray",
        "abbr": "TRAY",
        "name": "Traymore at Park Cities",
        "city": "Dallas, TX",
        "shootDate": "Tuesday, November 24, 2026",
        "shortDate": "11/24",
        "dayOfWeek": "Tuesday",
        "loungeName": "Traymore Highland Park Holiday Studio",
    },
    {
        "code": "mml",
        "abbr": "MML",
        "name": "Midland Medical Lodge",
        "city": "Midland, TX",
        "shootDate": "Monday, November 30, 2026",
        "shortDate": "11/30",
        "dayOfWeek": "Monday",
        "loungeName": "Midland Rose Garden Recreation Room",
    },
    {
        "code": "mmr",
        "abbr": "MMR",
        "name": "Madison Medical Resort",
        "city": "Odessa, TX",
        "shootDate": "Tuesday, December 1, 2026",
        "shortDate": "12/1",
        "dayOfWeek": "Tuesday",
        "loungeName": "Madison Grand Ballroom Staging Suite",
    },
    {
        "code": "aml",
        "abbr": "AML",
        "name": "Ashton Medical Lodge",
        "city": "Midland, TX",
        "shootDate": "Wednesday, December 2, 2026",
        "shortDate": "12/2",
        "dayOfWeek": "Wednesday",
        "loungeName": "Ashton Main Fireside Staging Lounge",
    },
    {
        "code": "sml",
        "abbr": "SML",
        "name": "Sheridan Medical Lodge",
        "city": "Burleson, TX",
        "shootDate": "Tuesday, December 8, 2026",
        "shortDate": "12/8",
        "dayOfWeek": "Tuesday",
        "loungeName": "Sheridan Chisholm Trail Gathering Room",
    },
    {
        "code": "scwf",
        "abbr": "SCWF",
        "name": "Senior Care Wichita Falls",
        "city": "Wichita Falls, TX",
        "shootDate": "Wednesday, December 9, 2026",
        "shortDate": "12/9",
        "dayOfWeek": "Wednesday",
        "loungeName": "Wichita Falls Red River Sunroom",
    },
]

# Standard 10-minute slots
MORNING_SLOTS = [
    '10:00 AM', '10:10 AM', '10:20 AM', '10:30 AM', '10:40 AM', '10:50 AM',
    '11:00 AM', '11:10 AM', '11:20 AM', '11:30 AM', '11:40 AM', '11:50 AM',
    '12:00 PM (LOCKED - Dept Head Picture)',
    '12:10 PM', '12:20 PM', '12:30 PM', '12:40 PM', '12:50 PM',
]
LUNCH_SLOT = ['01:00 PM - 01:45 PM (Photographer Lunch Break & Reset)']
AFTERNOON_SLOTS = [
    '01:45 PM', '01:55 PM',
    '02:05 PM', '02:15 PM', '02:25 PM', '02:35 PM', '02:45 PM', '02:55 PM',
    '03:05 PM', '03:15 PM', '03:25 PM', '03:35 PM', '03:45 PM', '03:55 PM',
    '04:05 PM', '04:15 PM', '04:25 PM', '04:35 PM', '04:45 PM', '04:55 PM',
    '05:05 PM', '05:15 PM', '05:25 PM', '05:35 PM', '05:45 PM', '05:55 PM',
    '06:05 PM', '06:15 PM', '06:25 PM', '06:35 PM', '06:45 PM', '06:55 PM',
]
ALL_SLOTS = MORNING_SLOTS + LUNCH_SLOT + AFTERNOON_SLOTS

HEADERS = [
    "Slot Time",
    "Status",
    "Booking Ref",
    "Resident Name",
    "Room #",
    "Family Contact",
    "Phone Number",
    "Email Address",
    "Mobility & Posing Needs",
    "Staff Notes / Internal",
    "Self-Service Reschedule Link",
    "Check-In Timestamp"
]

def style_worksheet(ws, facility, date_str):
    # Fonts
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="F8FAFC")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    locked_font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    break_font = Font(name="Calibri", size=10, bold=True, color="854D0E")
    regular_font = Font(name="Calibri", size=10, color="0F172A")
    time_font = Font(name="Calibri", size=10, bold=True, color="0F172A")

    # Fills
    title_fill = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid") # Deep Forest Green
    subtitle_fill = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
    header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid") # Holiday Red / Crimson
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    locked_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light red
    break_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # Light yellow

    # Borders
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # Title Row (Row 1)
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = f"🎄 FOURSQUARE HEALTHCARE 2026 CHRISTMAS PHOTO TOUR — {facility['name'].upper()} ({facility['abbr']})"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Subtitle Row (Row 2)
    ws.merge_cells("A2:L2")
    sub_cell = ws["A2"]
    sub_cell.value = f"📍 Location: {facility['city']}  |  📅 Date: {date_str}  |  🛋️ Staging Lounge: {facility['loungeName']}  |  🔗 Online Portal: https://foursquare-christmas-photoshoot.netlify.app/{facility['code']}"
    sub_cell.font = subtitle_font
    sub_cell.fill = subtitle_fill
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 26

    # Empty Row 3
    ws.row_dimensions[3].height = 10

    # Header Row (Row 4)
    ws.row_dimensions[4].height = 28
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows (Row 5+)
    current_row = 5
    for slot in ALL_SLOTS:
        is_locked_dept = "Dept Head Picture" in slot
        is_lunch = "Lunch Break" in slot

        status = "LOCKED" if is_locked_dept else ("BREAK" if is_lunch else "Open")
        notes = "Reserved exclusively for Facility Department Head Portrait" if is_locked_dept else ("Photographer meal break & backdrop lighting reset" if is_lunch else "")

        ws.row_dimensions[current_row].height = 22

        row_data = [
            slot,
            status,
            "", # Booking Ref
            "", # Resident Name
            "", # Room Number
            "", # Family Contact
            "", # Phone Number
            "", # Email Address
            "", # Mobility Needs
            notes, # Staff Notes
            f"https://foursquare-christmas-photoshoot.netlify.app/reschedule" if not is_locked_dept and not is_lunch else "",
            "" # Check-In Timestamp
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.border = thin_border

            if is_locked_dept:
                cell.fill = locked_fill
                cell.font = locked_font
            elif is_lunch:
                cell.fill = break_fill
                cell.font = break_font
            else:
                if current_row % 2 == 0:
                    cell.fill = zebra_fill
                cell.font = time_font if col_idx == 1 else regular_font

            if col_idx in [1, 2, 3, 5, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Freeze Panes below headers
    ws.freeze_panes = "A5"

    # Set Column Widths
    col_widths = {
        "A": 26, # Slot Time
        "B": 14, # Status
        "C": 15, # Booking Ref
        "D": 22, # Resident Name
        "E": 12, # Room #
        "F": 22, # Family Contact
        "G": 18, # Phone
        "H": 26, # Email
        "I": 30, # Mobility Needs
        "J": 35, # Staff Notes
        "K": 35, # Reschedule Link
        "L": 20  # Check-In Time
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

def generate_excel_files():
    base_dir = r"c:\Users\Corey\.gemini\antigravity-ide\scratch\foursquare-christmas-photoshoot\sheets"
    os.makedirs(base_dir, exist_ok=True)

    # 1. Master Workbook containing all facilities
    master_wb = openpyxl.Workbook()
    master_wb.remove(master_wb.active) # remove default sheet

    print("Generating Master Multi-Facility Workbook...")
    for fac in FACILITIES:
        if "dates" in fac and len(fac["dates"]) > 1:
            for d in fac["dates"]:
                ws = master_wb.create_sheet(title=f"{fac['abbr']} {d['shortDate'].replace('/', '-')}")
                style_worksheet(ws, fac, d["dateStr"])
        else:
            ws = master_wb.create_sheet(title=f"{fac['abbr']} ({fac['shortDate'].replace('/', '-')})")
            style_worksheet(ws, fac, fac["shootDate"])

    master_path = os.path.join(base_dir, "Foursquare_Christmas_Photo_Tour_2026_MASTER.xlsx")
    master_wb.save(master_path)
    print(f"[OK] Saved Master Workbook: {master_path}")

    # 2. Individual Workbooks for each facility
    print("\nGenerating Individual Facility Workbooks...")
    for fac in FACILITIES:
        wb = openpyxl.Workbook()
        if "dates" in fac and len(fac["dates"]) > 1:
            wb.remove(wb.active)
            for d in fac["dates"]:
                ws = wb.create_sheet(title=d["label"])
                style_worksheet(ws, fac, d["dateStr"])
        else:
            ws = wb.active
            ws.title = "Schedule & Run-Sheet"
            style_worksheet(ws, fac, fac["shootDate"])

        filename = f"{fac['abbr']}_{fac['name'].replace(' ', '_').replace('&', 'and')}_Photo_Tour_2026.xlsx"
        path = os.path.join(base_dir, filename)
        wb.save(path)
        print(f"[OK] Saved: {filename}")

    print("\nAll Excel workbooks generated successfully!")

def generate_apps_script_file():
    apps_script_code = """/**
 * ============================================================================
 * 🎄 FOURSQUARE HEALTHCARE 2026 CHRISTMAS PHOTO TOUR
 * Google Sheets Auto-Provisioning & Sharing Script
 *
 * INSTRUCTIONS:
 * 1. Open https://script.new in your browser (logged into your Thinkments/Google account)
 * 2. Delete any boilerplate code and paste this entire file.
 * 3. Click "Run" -> "createAllFoursquareFacilitySheets".
 * 4. Google will ask for permission once; click "Review permissions" -> select account -> "Allow".
 * 5. In under 15 seconds, all 12 facility sheets will be created, styled, and shared with jeremy@thinkments.com!
 * ============================================================================
 */

function createAllFoursquareFacilitySheets() {
  const SHARE_EMAIL = "jeremy@thinkments.com";
  
  const FACILITIES = [
    {
      code: "hml",
      abbr: "HML",
      name: "Hillside Medical Lodge",
      city: "Beeville, TX",
      shootDate: "Thursday, November 5, 2026",
      shortDate: "11/5",
      loungeName: "Hillside Fireside Solarium"
    },
    {
      code: "wnr",
      abbr: "WNR",
      name: "Whitney Nursing & Rehabilitation",
      city: "Whitney, TX",
      shootDate: "Friday, November 6, 2026",
      shortDate: "11/6",
      loungeName: "Whitney Heritage Community Room"
    },
    {
      code: "cml",
      abbr: "CML",
      name: "Cheyenne Medical Lodge",
      city: "Colorado City, TX",
      shootDate: "Monday, November 9 & Tuesday, November 10, 2026",
      shortDate: "11/9 & 11/10",
      dates: [
        { label: "Day 1 (Nov 9)", dateStr: "Monday, November 9, 2026" },
        { label: "Day 2 (Nov 10)", dateStr: "Tuesday, November 10, 2026" }
      ],
      loungeName: "Cheyenne Grand Prairie Staging Room"
    },
    {
      code: "pml",
      abbr: "PML",
      name: "Princeton Medical Lodge",
      city: "Princeton, TX",
      shootDate: "Thursday, November 12, 2026",
      shortDate: "11/12",
      loungeName: "Princeton Courtyard Pavilion"
    },
    {
      code: "fhr",
      abbr: "FHR",
      name: "Farmersville Health & Rehabilitation",
      city: "Farmersville, TX",
      shootDate: "Friday, November 13, 2026",
      shortDate: "11/13",
      loungeName: "Farmersville Evergreen Great Room"
    },
    {
      code: "lml",
      abbr: "LML",
      name: "Lexington Medical Lodge",
      city: "Farmersville, TX",
      shootDate: "Tuesday, November 17, 2026",
      shortDate: "11/17",
      loungeName: "Lexington Magnolia Activity Atrium"
    },
    {
      code: "tray",
      abbr: "TRAY",
      name: "Traymore at Park Cities",
      city: "Dallas, TX",
      shootDate: "Tuesday, November 24, 2026",
      shortDate: "11/24",
      loungeName: "Traymore Highland Park Holiday Studio"
    },
    {
      code: "mml",
      abbr: "MML",
      name: "Midland Medical Lodge",
      city: "Midland, TX",
      shootDate: "Monday, November 30, 2026",
      shortDate: "11/30",
      loungeName: "Midland Rose Garden Recreation Room"
    },
    {
      code: "mmr",
      abbr: "MMR",
      name: "Madison Medical Resort",
      city: "Odessa, TX",
      shootDate: "Tuesday, December 1, 2026",
      shortDate: "12/1",
      loungeName: "Madison Grand Ballroom Staging Suite"
    },
    {
      code: "aml",
      abbr: "AML",
      name: "Ashton Medical Lodge",
      city: "Midland, TX",
      shootDate: "Wednesday, December 2, 2026",
      shortDate: "12/2",
      loungeName: "Ashton Main Fireside Staging Lounge"
    },
    {
      code: "sml",
      abbr: "SML",
      name: "Sheridan Medical Lodge",
      city: "Burleson, TX",
      shootDate: "Tuesday, December 8, 2026",
      shortDate: "12/8",
      loungeName: "Sheridan Chisholm Trail Gathering Room"
    },
    {
      code: "scwf",
      abbr: "SCWF",
      name: "Senior Care Wichita Falls",
      city: "Wichita Falls, TX",
      shootDate: "Wednesday, December 9, 2026",
      shortDate: "12/9",
      loungeName: "Wichita Falls Red River Sunroom"
    }
  ];

  const MORNING_SLOTS = [
    '10:00 AM', '10:10 AM', '10:20 AM', '10:30 AM', '10:40 AM', '10:50 AM',
    '11:00 AM', '11:10 AM', '11:20 AM', '11:30 AM', '11:40 AM', '11:50 AM',
    '12:00 PM (LOCKED - Dept Head Picture)',
    '12:10 PM', '12:20 PM', '12:30 PM', '12:40 PM', '12:50 PM',
  ];
  const LUNCH_SLOT = ['01:00 PM - 01:45 PM (Photographer Lunch Break & Reset)'];
  const AFTERNOON_SLOTS = [
    '01:45 PM', '01:55 PM',
    '02:05 PM', '02:15 PM', '02:25 PM', '02:35 PM', '02:45 PM', '02:55 PM',
    '03:05 PM', '03:15 PM', '03:25 PM', '03:35 PM', '03:45 PM', '03:55 PM',
    '04:05 PM', '04:15 PM', '04:25 PM', '04:35 PM', '04:45 PM', '04:55 PM',
    '05:05 PM', '05:15 PM', '05:25 PM', '05:35 PM', '05:45 PM', '05:55 PM',
    '06:05 PM', '06:15 PM', '06:25 PM', '06:35 PM', '06:45 PM', '06:55 PM',
  ];
  const ALL_SLOTS = MORNING_SLOTS.concat(LUNCH_SLOT).concat(AFTERNOON_SLOTS);

  const HEADERS = [
    "Slot Time",
    "Status",
    "Booking Ref",
    "Resident Name",
    "Room #",
    "Family Contact",
    "Phone Number",
    "Email Address",
    "Mobility & Special Posing Needs",
    "Staff Notes / Internal",
    "Self-Service Reschedule Link",
    "Check-In Timestamp"
  ];

  Logger.log("🚀 Starting Foursquare Google Sheet Generation...");

  const createdSheets = [];

  // Create a dedicated Google Drive folder for organization
  const folderName = "Foursquare Christmas Photo Tour 2026 - Facility Sheets";
  let targetFolder;
  const existingFolders = DriveApp.getFoldersByName(folderName);
  if (existingFolders.hasNext()) {
    targetFolder = existingFolders.next();
  } else {
    targetFolder = DriveApp.createFolder(folderName);
  }
  
  // Share the entire folder with Jeremy
  try {
    targetFolder.addEditor(SHARE_EMAIL);
    Logger.log("✓ Shared master folder with: " + SHARE_EMAIL);
  } catch (e) {
    Logger.log("Folder share note: " + e.message);
  }

  FACILITIES.forEach(function(fac, idx) {
    const title = "Foursquare Photo Tour 2026 - " + fac.name + " (" + fac.abbr + ")";
    Logger.log("[" + (idx + 1) + "/" + FACILITIES.length + "] Creating: " + title);

    const ss = SpreadsheetApp.create(title);
    const driveFile = DriveApp.getFileById(ss.getId());
    
    // Move to dedicated folder
    driveFile.moveTo(targetFolder);

    // Share each file with Jeremy
    try {
      driveFile.addEditor(SHARE_EMAIL);
      Logger.log("   ✓ Shared with " + SHARE_EMAIL);
    } catch(err) {
      Logger.log("   ⚠️ Share note: " + err.message);
    }

    if (fac.dates && fac.dates.length > 1) {
      // Multi-day facility
      fac.dates.forEach(function(d, dIdx) {
        let ws;
        if (dIdx === 0) {
          ws = ss.getActiveSheet();
          ws.setName(d.label);
        } else {
          ws = ss.insertSheet(d.label);
        }
        formatSheet(ws, fac, d.dateStr, ALL_SLOTS, HEADERS);
      });
    } else {
      const ws = ss.getActiveSheet();
      ws.setName("Schedule & Roster");
      formatSheet(ws, fac, fac.shootDate, ALL_SLOTS, HEADERS);
    }

    createdSheets.push({
      facility: fac.name,
      abbr: fac.abbr,
      code: fac.code,
      url: ss.getUrl()
    });

    Logger.log("   ✓ Created URL: " + ss.getUrl());
  });

  Logger.log("\\n=================================================================");
  Logger.log("🎉 ALL 12 FOURSQUARE SHEETS CREATED & SHARED SUCCESSFULLY!");
  Logger.log("📂 Drive Folder: " + targetFolder.getUrl());
  Logger.log("=================================================================\\n");

  createdSheets.forEach(function(item) {
    Logger.log(item.abbr + " (" + item.facility + "): " + item.url);
  });
}

function formatSheet(ws, fac, dateStr, slots, headers) {
  // Clear any existing grid
  ws.clear();

  // Set up Title (Row 1)
  ws.getRange("A1:L1").merge();
  const titleCell = ws.getRange("A1");
  titleCell.setValue("🎄 FOURSQUARE HEALTHCARE 2026 CHRISTMAS PHOTO TOUR — " + fac.name.toUpperCase() + " (" + fac.abbr + ")");
  titleCell.setBackground("#1B4332"); // Deep Forest Green
  titleCell.setFontColor("#FFFFFF");
  titleCell.setFontSize(14);
  titleCell.setFontWeight("bold");
  titleCell.setHorizontalAlignment("center");
  titleCell.setVerticalAlignment("middle");
  ws.setRowHeight(1, 40);

  // Subtitle (Row 2)
  ws.getRange("A2:L2").merge();
  const subCell = ws.getRange("A2");
  subCell.setValue("📍 Location: " + fac.city + "  |  📅 Date: " + dateStr + "  |  🛋️ Lounge: " + fac.loungeName + "  |  🔗 https://foursquare-christmas-photoshoot.netlify.app/" + fac.code);
  subCell.setBackground("#2D6A4F");
  subCell.setFontColor("#F8FAFC");
  subCell.setFontSize(10);
  subCell.setFontWeight("bold");
  subCell.setHorizontalAlignment("center");
  subCell.setVerticalAlignment("middle");
  ws.setRowHeight(2, 26);

  // Headers (Row 4)
  ws.setRowHeight(3, 10);
  ws.setRowHeight(4, 28);
  const headerRange = ws.getRange(4, 1, 1, headers.length);
  headerRange.setValues([headers]);
  headerRange.setBackground("#8B0000"); // Holiday Red / Burgundy
  headerRange.setFontColor("#FFFFFF");
  headerRange.setFontWeight("bold");
  headerRange.setFontSize(10);
  headerRange.setHorizontalAlignment("center");
  headerRange.setVerticalAlignment("middle");

  // Populate slots
  const rows = [];
  slots.forEach(function(slot) {
    const isLocked = slot.indexOf("Dept Head Picture") !== -1;
    const isBreak = slot.indexOf("Lunch Break") !== -1;
    const status = isLocked ? "LOCKED" : (isBreak ? "BREAK" : "Open");
    const notes = isLocked ? "Reserved exclusively for Facility Department Head Portrait" : (isBreak ? "Photographer meal break & studio lighting reset" : "");
    const resched = (!isLocked && !isBreak) ? ("https://foursquare-christmas-photoshoot.netlify.app/reschedule") : "";

    rows.push([
      slot,
      status,
      "", // Booking Ref
      "", // Resident Name
      "", // Room #
      "", // Family Contact
      "", // Phone
      "", // Email
      "", // Mobility
      notes,
      resched,
      ""  // Check-In
    ]);
  });

  const dataRange = ws.getRange(5, 1, rows.length, headers.length);
  dataRange.setValues(rows);
  dataRange.setFontSize(9);
  dataRange.setVerticalAlignment("middle");

  // Conditional background styling for locked & break rows
  for (let i = 0; i < slots.length; i++) {
    const rowNum = 5 + i;
    const slotText = slots[i];
    if (slotText.indexOf("Dept Head Picture") !== -1) {
      ws.getRange(rowNum, 1, 1, headers.length).setBackground("#FEE2E2").setFontColor("#991B1B").setFontWeight("bold");
    } else if (slotText.indexOf("Lunch Break") !== -1) {
      ws.getRange(rowNum, 1, 1, headers.length).setBackground("#FEF9C3").setFontColor("#854D0E").setFontWeight("bold");
    } else if (i % 2 === 1) {
      ws.getRange(rowNum, 1, 1, headers.length).setBackground("#F8FAFC");
    }
  }

  // Freeze panes
  ws.setFrozenRows(4);

  // Column Widths
  const widths = [190, 90, 100, 150, 80, 150, 120, 180, 200, 240, 240, 130];
  widths.forEach(function(w, i) {
    ws.setColumnWidth(i + 1, w);
  });
}
"""
    script_path = r"c:\Users\Corey\.gemini\antigravity-ide\scratch\foursquare-christmas-photoshoot\scripts\GoogleAppsScript_CreateSheets.js"
    with open(script_path, "w", encoding="utf-8") as f:
        f.write(apps_script_code)
    print(f"[OK] Saved Google Apps Script: {script_path}")

if __name__ == "__main__":
    generate_excel_files()
    generate_apps_script_file()
